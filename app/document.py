"""Word document builder (python-docx) and post-save verification, plus heading matching."""
import difflib
import os
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor

from .schemas import Assumption, DocumentSection, DocumentTable

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
FILENAME_RE = re.compile(r"^[0-9a-f]{32}\.docx$")

# Sections the system always appends itself. The drafter is told not to write them.
KNOWN_FACTS_HEADING = "Known Facts (Provided in the Request)"
ASSUMPTIONS_HEADING = "Assumptions and Open Questions"
DISCLOSURE_HEADING = "AI-Assisted Generation Disclosure and Human-Review Notice"
SYSTEM_HEADINGS = (KNOWN_FACTS_HEADING, ASSUMPTIONS_HEADING, DISCLOSURE_HEADING)

CATEGORY_LABELS = {
    "assumption": "Assumption",
    "illustrative_data": "Illustrative data",
    "requires_confirmation": "Requires confirmation",
    "constraint_conflict": "Constraint conflict",
}


def get_output_dir() -> Path:
    base = os.getenv("OUTPUT_DIR")
    path = Path(base) if base else Path(__file__).resolve().parent.parent / "outputs"
    path.mkdir(parents=True, exist_ok=True)
    return path.resolve()


# ---------- heading matching (deterministic section checks) ----------

def normalize_heading(text: str) -> str:
    text = text.lower().replace("&", " and ")
    text = re.sub(r"^\s*\d+[\.\)]?\s+", "", text)  # drop leading "1." / "2)"
    text = re.sub(r"[^a-z0-9 ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


FUZZY_MIN_CHARS = 12
FUZZY_THRESHOLD = 0.85


def _stem(token: str) -> str:
    """Treat simple plurals as the same word ('risks' = 'risk'); 'ss' endings are left alone."""
    return token[:-1] if len(token) > 3 and token.endswith("s") and not token.endswith("ss") else token


def _canonical(text: str) -> str:
    return " ".join(_stem(token) for token in normalize_heading(text).split())


def heading_matches(a: str, b: str) -> bool:
    """Do two headings mean the same section?

    1. Equal after normalisation (case, numbering, '&' vs 'and', punctuation, simple plurals).
    2. Otherwise a bounded fuzzy match, ONLY when both headings have 2+ words and 12+ characters,
       at similarity >= 0.85 (typos and small wording variants such as 'Analysis' / 'Analyses').
    There is deliberately NO substring rule: 'Scope' must not match 'Microscope Details', and a bare
    'Risk' must not match 'Risk Management'."""
    x, y = _canonical(a), _canonical(b)
    if not x or not y:
        return False
    if x == y:
        return True
    if len(x.split()) < 2 or len(y.split()) < 2 or min(len(x), len(y)) < FUZZY_MIN_CHARS:
        return False
    return difflib.SequenceMatcher(None, x, y).ratio() >= FUZZY_THRESHOLD


# Names the planner may use for the sections Python renders itself. An explicit list, not substrings.
SYSTEM_ALIASES = {
    _canonical(name)
    for name in (
        *SYSTEM_HEADINGS,
        "Assumptions", "Assumptions Register", "Open Questions", "Open Questions and Assumptions",
        "Known Facts", "Facts Provided in the Request",
        "Disclosure", "AI Disclosure", "Human Review Notice", "Disclosure and Human Review Notice",
        "AI-Assisted Generation Disclosure",
    )
}


def is_system_section(name: str) -> bool:
    return _canonical(name) in SYSTEM_ALIASES or any(heading_matches(name, h) for h in SYSTEM_HEADINGS)


def missing_sections(required: list[str], headings: list[str]) -> list[str]:
    """Required sections with no matching heading. System-owned sections are always rendered by
    Python, so they are never reported missing."""
    return [
        name for name in required
        if not is_system_section(name) and not any(heading_matches(name, h) for h in headings)
    ]


# ---------- building ----------

@dataclass
class DocumentSpec:
    run_id: str
    title: str
    document_type: str
    audience: str
    sections: list[DocumentSection]
    known_facts: list[str] = field(default_factory=list)
    assumptions: list[Assumption] = field(default_factory=list)
    open_questions: list[str] = field(default_factory=list)
    remaining_issues: list[str] = field(default_factory=list)
    fallback_notice: str | None = None  # set only for fallback documents; states what really happened
    provider: str = ""
    model: str = ""


def _clean(text: str) -> str:
    """Strip stray Markdown the model may emit despite instructions."""
    text = text.replace("**", "").replace("__", "").replace("`", "")
    return re.sub(r"^\s{0,3}#{1,6}\s+", "", text, flags=re.MULTILINE).strip()


def _style(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for name in ("Title", "Heading 1"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(1)


def _add_table(doc: Document, table: DocumentTable) -> None:
    grid = doc.add_table(rows=1, cols=len(table.columns))
    try:
        grid.style = "Light Grid Accent 1"
    except KeyError:
        grid.style = "Table Grid"
    for i, column in enumerate(table.columns):
        cell = grid.rows[0].cells[i]
        cell.text = ""
        cell.paragraphs[0].add_run(_clean(column)).bold = True
    for row in table.rows:
        cells = grid.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = _clean(value)
    doc.add_paragraph()


def _add_section(doc: Document, section: DocumentSection) -> None:
    doc.add_heading(_clean(section.heading), level=1)
    for paragraph in re.split(r"\n\s*\n", section.content.strip()):
        if paragraph.strip():
            doc.add_paragraph(_clean(paragraph))
    for bullet in section.bullets:
        doc.add_paragraph(_clean(bullet), style="List Bullet")
    if section.table and section.table.rows:
        _add_table(doc, section.table)


def build_docx(spec: DocumentSpec) -> Path:
    """Write <run_id>.docx into the output directory (server-generated filename only)."""
    doc = Document()
    _style(doc)
    doc.add_heading(_clean(spec.title), level=0)
    doc.add_paragraph(f"{_clean(spec.document_type).title()}  |  Prepared for: {_clean(spec.audience)}")
    stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    doc.add_paragraph(f"AI-assisted draft  |  Run ID: {spec.run_id}  |  Generated: {stamp}")

    if spec.fallback_notice:
        notice = doc.add_paragraph()
        run = notice.add_run(f"FALLBACK DOCUMENT: {spec.fallback_notice}")
        run.bold = True
        run.font.color.rgb = RGBColor(0xB0, 0x1E, 0x1E)

    for section in spec.sections:
        _add_section(doc, section)

    doc.add_heading(KNOWN_FACTS_HEADING, level=1)
    if spec.known_facts:
        for fact in spec.known_facts:
            doc.add_paragraph(_clean(fact), style="List Bullet")
    else:
        doc.add_paragraph("None recorded.")

    doc.add_heading(ASSUMPTIONS_HEADING, level=1)
    if spec.assumptions:
        _add_table(doc, DocumentTable(
            columns=["Type", "Statement", "Impact"],
            rows=[[CATEGORY_LABELS.get(a.category, a.category), a.statement, a.impact] for a in spec.assumptions],
        ))
    else:
        doc.add_paragraph("No assumptions were recorded.")
    doc.add_paragraph("Open questions", style="Heading 2")
    for question in spec.open_questions or ["None recorded."]:
        doc.add_paragraph(_clean(question), style="List Bullet")
    if spec.remaining_issues:
        doc.add_paragraph("Unresolved review findings", style="Heading 2")
        for issue in spec.remaining_issues:
            doc.add_paragraph(_clean(issue), style="List Bullet")

    doc.add_heading(DISCLOSURE_HEADING, level=1)
    doc.add_paragraph(
        f"This document was generated with AI assistance (provider: {spec.provider}, model: {spec.model}). "
        "Content, including any figures marked as illustrative, has not been independently verified."
    )
    doc.add_paragraph(
        "Decisions with legal, financial, security, operational or other significant consequences "
        "require review and approval by a qualified person before this document is relied upon."
    )

    path = get_output_dir() / f"{spec.run_id}.docx"
    doc.save(path)
    return path


# ---------- verification ----------

@dataclass
class VerificationResult:
    ok: bool
    problems: list[str] = field(default_factory=list)
    missing_sections: list[str] = field(default_factory=list)
    headings: list[str] = field(default_factory=list)


def verify_docx(path: Path, expected_title: str, required_sections: list[str]) -> VerificationResult:
    """Reopen the saved file. File integrity problems fail the result; missing section headings
    are reported separately so they can be disclosed as warnings."""
    problems: list[str] = []
    output_dir = get_output_dir()
    resolved = Path(path).resolve()
    if resolved.parent != output_dir:
        return VerificationResult(False, ["File is outside the output directory"])
    if not resolved.is_file() or resolved.stat().st_size == 0:
        return VerificationResult(False, ["File is missing or empty"])
    try:
        doc = Document(str(resolved))
    except Exception as exc:  # corrupt or non-docx file
        return VerificationResult(False, [f"File could not be reopened ({type(exc).__name__})"])

    titles = [p.text for p in doc.paragraphs if p.style.name == "Title"]
    if _clean(expected_title) not in titles:
        problems.append("Title not found in document")
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    return VerificationResult(
        ok=not problems,
        problems=problems,
        missing_sections=missing_sections(required_sections, headings),
        headings=headings,
    )
