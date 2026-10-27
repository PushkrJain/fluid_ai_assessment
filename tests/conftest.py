import json

import pytest
from fastapi.testclient import TestClient

from app import llm
from app.main import app

STAGES = ("planner", "analyzer", "drafter", "reviewer", "reviser", "repair")


class FakeLLM:
    """Replaces llm.generate_text. Responses are queued per pipeline stage; each call pops one."""

    def __init__(self):
        self.queues: dict[str, list] = {}
        self.calls: list[str] = []
        self.users: list[tuple[str, str]] = []
        self.prompts = {name: llm.load_prompt(name) for name in STAGES}

    def script(self, stage: str, *items) -> "FakeLLM":
        self.queues.setdefault(stage, []).extend(items)
        return self

    def count(self, stage: str) -> int:
        return self.calls.count(stage)

    def context(self, stage: str, index: int = -1) -> dict:
        """Parse the JSON context the given stage received (for asserting on prompts' inputs)."""
        messages = [u for st, u in self.users if st == stage]
        body = messages[index].split("<context>\n", 1)[1].rsplit("\n</context>", 1)[0]
        return json.loads(body)

    def __call__(self, system, user, **kwargs):
        stage = next((n for n, p in self.prompts.items() if p == system), "unknown")
        self.calls.append(stage)
        self.users.append((stage, user))
        queue = self.queues.get(stage)
        if not queue:
            raise AssertionError(f"Unexpected LLM call for stage '{stage}'")
        item = queue.pop(0)
        if isinstance(item, Exception):
            raise item
        return item if isinstance(item, str) else json.dumps(item)


def task(id_, handler, deps=()):
    return {"id": id_, "title": f"Task {id_}", "description": f"Do {handler}", "handler": handler,
            "dependencies": list(deps), "expected_output": "Output"}


def make_plan(**over):
    plan = {
        "document_type": "project plan", "title": "Test Plan", "audience": "Leaders",
        "known_facts": ["Eight weeks"], "requirements": ["Cover scope and risks"], "constraints": ["Eight weeks"],
        "assumptions": [{"statement": "Team is full time", "category": "assumption", "impact": "Schedule"}],
        "open_questions": ["Who owns launch?"],
        "required_sections": ["Executive Summary", "Scope", "Risks and Mitigations", "Assumptions and Open Questions"],
        "tasks": [task(1, "analyze_request"), task(2, "resolve_assumptions", [1]), task(3, "draft_document", [2]),
                  task(4, "review_document", [3]), task(5, "generate_docx", [4])],
    }
    plan.update(over)
    return plan


def make_analysis(**over):
    note = {"summary": "Looks feasible with care.", "findings": ["f1"], "conflicts": ["Scope vs time may conflict"],
            "open_questions": ["Which department pilots first?"]}
    note.update(over)
    return note


def section(heading, **over):
    data = {"heading": heading, "content": f"Content for {heading}.", "bullets": []}
    data.update(over)
    return data


def make_draft(headings=("Executive Summary", "Scope", "Risks and Mitigations")):
    return {"title": "Test Plan", "sections": [
        section(h, table={"columns": ["Item", "Detail"], "rows": [["a", 1]]}) if h.startswith("Risks") else section(h)
        for h in headings]}


def make_review(**over):
    report = {"score": 88, "identified_gaps": [], "failed_checks": [], "sections_to_revise": [],
              "recommended_corrections": [], "remaining_issues": [], "requirements_covered": ["Scope"]}
    report.update(over)
    return report


@pytest.fixture
def fake_llm(monkeypatch, tmp_path):
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-key")
    monkeypatch.setenv("OPENROUTER_MODEL", "test/model")
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path))
    fake = FakeLLM()
    monkeypatch.setattr(llm, "generate_text", fake)
    return fake


def script_happy(fake):
    """Queue one clean run: one call per stage, no revision."""
    fake.script("planner", make_plan()).script("analyzer", make_analysis())
    fake.script("drafter", make_draft()).script("reviewer", make_review())
    return fake


@pytest.fixture
def happy(fake_llm):
    return script_happy(fake_llm)


@pytest.fixture
def client():
    return TestClient(app)


REQUEST = {"request": "Create a project plan for launching an onboarding portal in eight weeks."}


# ---------- shared helpers ----------

def docx_text(path) -> str:
    from docx import Document
    doc = Document(str(path))
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return "\n".join([p.text for p in doc.paragraphs] + cells)


def docx_headings(path) -> list[str]:
    from docx import Document
    return [p.text for p in Document(str(path)).paragraphs if p.style.name.startswith("Heading")]


def make_state(headings=("Executive Summary", "Scope", "Risks and Mitigations"), **plan_over):
    """An AgentState with a validated plan and a draft, for calling handlers directly."""
    from app.agent import AgentState
    from app.schemas import DocumentContent, Plan, Task
    state = AgentState(request="Create a project plan for an onboarding portal.", provider="openrouter", model="test/model")
    state.plan = Plan.model_validate(make_plan(**plan_over))
    state.tasks = [Task.from_planned(t) for t in state.plan.tasks]
    state.assumptions = list(state.plan.assumptions)
    state.open_questions = list(state.plan.open_questions)
    state.draft = DocumentContent.model_validate(make_draft(headings))
    return state


class FakeResponse:
    def __init__(self, text, status="completed"):
        self.output_text, self.status = text, status


class RoutingClient:
    """Stands in for the OpenAI client so the REAL llm.generate_text path (logging, retry) runs."""

    def __init__(self, fake):
        self.fake, self.responses = fake, self

    def create(self, *, model, input, max_output_tokens, **kwargs):
        return FakeResponse(self.fake(input[0]["content"], input[1]["content"]))


@pytest.fixture
def fake_client(monkeypatch, tmp_path):
    monkeypatch.setenv("OPENROUTER_API_KEY", "sk-test-secret-key-123456")
    monkeypatch.setenv("OPENROUTER_MODEL", "test/model")
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path))
    monkeypatch.setattr(llm.time, "sleep", lambda s: None)
    fake = FakeLLM()
    monkeypatch.setattr(llm, "_make_client", lambda settings: RoutingClient(fake))
    return fake
