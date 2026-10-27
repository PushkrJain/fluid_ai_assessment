import pytest
from docx import Document

from app import llm
from app.handlers import HANDLERS
from app.schemas import HandlerName
from tests.conftest import (REQUEST, make_analysis, make_draft, make_plan, make_review, section, task)
from typing import get_args


def headings(path):
    return [p.text for p in Document(str(path)).paragraphs if p.style.name.startswith("Heading")]


def test_happy_path_runs_every_planned_task(happy, client, tmp_path):
    r = client.post("/agent", json=REQUEST)
    assert r.status_code == 200
    body = r.json()
    assert body["status"] in ("completed", "completed_with_warnings")
    assert [t["status"] for t in body["plan"]] == ["completed"] * 5
    assert body["fallback_used"] is False and body["revisions_performed"] == 0
    assert body["provider"] == "openrouter" and body["model"] == "test/model"
    # analysis conflict was folded into the register and the open questions by resolve_assumptions
    assert any(a["category"] == "constraint_conflict" for a in body["assumptions"])
    assert "Which department pilots first?" in body["open_questions"]
    assert body["review"]["score"] == 88
    # real file, served through the real endpoint, reopens with planner-declared sections
    download = client.get(body["download_url"])
    assert download.status_code == 200
    path = tmp_path / body["document_filename"]
    assert {"Executive Summary", "Scope", "Risks and Mitigations", "Assumptions and Open Questions"} <= set(headings(path))
    assert Document(str(path)).tables, "risk table should be rendered"


def test_plan_drives_execution(happy, client):
    """A plan without the optional steps must run without them (no hard-coded pipeline)."""
    happy.queues["planner"] = [make_plan(tasks=[task(1, "draft_document"), task(2, "review_document", [1]),
                                                task(3, "generate_docx", [2])])]
    body = client.post("/agent", json=REQUEST).json()
    assert [t["handler"] for t in body["plan"]] == ["draft_document", "review_document", "generate_docx"]
    assert happy.count("analyzer") == 0


def test_short_or_blank_request_is_422(client):
    assert client.post("/agent", json={"request": "too short"}).status_code == 422
    assert client.post("/agent", json={"request": "          "}).status_code == 422
    assert client.post("/agent", json={}).status_code == 422


def test_malformed_plan_gets_exactly_one_repair(fake_llm, client):
    fake_llm.script("planner", "not json at all").script("repair", make_plan())
    fake_llm.script("analyzer", make_analysis()).script("drafter", make_draft()).script("reviewer", make_review())
    body = client.post("/agent", json=REQUEST).json()
    assert body["fallback_used"] is False and fake_llm.count("repair") == 1


def test_unusable_plan_after_repair_uses_labelled_fallback(fake_llm, client, tmp_path):
    fake_llm.script("planner", "garbage").script("repair", "still garbage")
    body = client.post("/agent", json=REQUEST).json()
    assert body["status"] == "completed_with_warnings" and body["fallback_used"] is True
    assert "No AI planning" in body["summary"]
    assert body["review"] is None
    assert [t["handler"] for t in body["plan"]] == ["generate_docx"]
    text = " ".join(p.text for p in Document(str(tmp_path / body["document_filename"])).paragraphs)
    assert "FALLBACK DOCUMENT" in text


def test_llm_outage_during_planning_uses_fallback(fake_llm, client):
    fake_llm.script("planner", llm.LLMError("LLM unavailable after 2 attempts"))
    body = client.post("/agent", json=REQUEST).json()
    assert body["fallback_used"] is True and body["download_url"]


def test_missing_configuration_returns_503_without_fallback(monkeypatch, tmp_path, client):
    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path))
    r = client.post("/agent", json=REQUEST)
    assert r.status_code == 503
    body = r.json()
    assert body["status"] == "failed" and body["download_url"] is None and "not configured" in body["error"]


def test_draft_failure_keeps_plan_but_labels_fallback(fake_llm, client):
    fake_llm.script("planner", make_plan()).script("analyzer", make_analysis())
    fake_llm.script("drafter", llm.LLMError("LLM unavailable"))
    body = client.post("/agent", json=REQUEST).json()
    statuses = {t["handler"]: t["status"] for t in body["plan"]}
    assert statuses == {"analyze_request": "completed", "resolve_assumptions": "completed",
                        "draft_document": "failed", "review_document": "skipped", "generate_docx": "completed"}
    assert body["fallback_used"] is True and body["status"] == "completed_with_warnings"


def test_noncritical_failure_is_recorded_and_pipeline_continues(fake_llm, client):
    fake_llm.script("planner", make_plan()).script("analyzer", llm.LLMError("LLM unavailable"))
    fake_llm.script("drafter", make_draft()).script("reviewer", make_review())
    body = client.post("/agent", json=REQUEST).json()
    statuses = {t["handler"]: t["status"] for t in body["plan"]}
    assert statuses["analyze_request"] == "failed" and statuses["generate_docx"] == "completed"
    assert body["status"] == "completed_with_warnings" and body["fallback_used"] is False
    assert any("analyze_request" in w for w in body["warnings"])


def test_review_failure_delivers_document_with_warning(fake_llm, client):
    fake_llm.script("planner", make_plan()).script("analyzer", make_analysis()).script("drafter", make_draft())
    fake_llm.script("reviewer", llm.LLMError("LLM unavailable"))
    body = client.post("/agent", json=REQUEST).json()
    assert body["review"] is None and body["download_url"] and body["status"] == "completed_with_warnings"


def test_known_gap_triggers_one_targeted_revision(fake_llm, client, tmp_path):
    """Draft omits a required section; reviewer flags it; exactly one revision adds it."""
    fake_llm.script("planner", make_plan()).script("analyzer", make_analysis())
    fake_llm.script("drafter", make_draft(headings=("Executive Summary", "Scope")))
    fake_llm.script("reviewer",
                    make_review(score=60, identified_gaps=["No risks section"], failed_checks=["Required section present"],
                                sections_to_revise=["Risks and Mitigations"]),
                    make_review(score=90))
    fake_llm.script("reviser", {"summary": "Added the risks section.", "sections": [section("Risks and Mitigations")]})
    body = client.post("/agent", json=REQUEST).json()
    assert body["revisions_performed"] == 1
    assert fake_llm.count("reviewer") == 2 and fake_llm.count("reviser") == 1
    assert body["review"]["score"] == 90  # the FINAL check is what gets reported
    assert "Risks and Mitigations" in headings(tmp_path / body["document_filename"])


def test_deterministic_check_catches_gap_the_reviewer_missed(fake_llm, client):
    fake_llm.script("planner", make_plan()).script("analyzer", make_analysis())
    fake_llm.script("drafter", make_draft(headings=("Executive Summary", "Scope")))
    fake_llm.script("reviewer", make_review(score=95), make_review(score=95))  # model claims all is well
    fake_llm.script("reviser", {"summary": "Added.", "sections": [section("Risks and Mitigations")]})
    body = client.post("/agent", json=REQUEST).json()
    assert body["revisions_performed"] == 1 and fake_llm.count("reviser") == 1


def test_revision_never_exceeds_one_round(fake_llm, client):
    bad = make_review(score=40, failed_checks=["Coverage"], sections_to_revise=["Scope"], remaining_issues=["Still weak"])
    fake_llm.script("planner", make_plan()).script("analyzer", make_analysis()).script("drafter", make_draft())
    fake_llm.script("reviewer", bad, bad, bad, bad)
    fake_llm.script("reviser", {"summary": "Tried.", "sections": [section("Scope")]},
                    {"summary": "Again.", "sections": [section("Scope")]})
    body = client.post("/agent", json=REQUEST).json()
    assert body["revisions_performed"] == 1
    assert fake_llm.count("reviewer") == 2 and fake_llm.count("reviser") == 1
    assert body["status"] == "completed_with_warnings" and "Still weak" in body["review"]["remaining_issues"]


def test_failed_revision_keeps_original_draft_and_discloses(fake_llm, client):
    fake_llm.script("planner", make_plan()).script("analyzer", make_analysis()).script("drafter", make_draft())
    fake_llm.script("reviewer", make_review(score=50, failed_checks=["x"], sections_to_revise=["Scope"],
                                            identified_gaps=["Scope is thin"]))
    fake_llm.script("reviser", llm.LLMError("LLM unavailable"))
    body = client.post("/agent", json=REQUEST).json()
    assert body["revisions_performed"] == 0 and body["download_url"]
    assert "Scope is thin" in body["review"]["remaining_issues"]
    assert any("Revision could not be completed" in w for w in body["warnings"])


def test_unexpected_handler_exception_is_not_leaked(happy, client, monkeypatch):
    def boom(state, task):
        raise ValueError("secret internal detail")
    monkeypatch.setitem(HANDLERS, "analyze_request", boom)
    body = client.post("/agent", json=REQUEST).json()
    failed = next(t for t in body["plan"] if t["handler"] == "analyze_request")
    assert failed["status"] == "failed" and "secret" not in failed["error"]


def test_registry_matches_schema_allowlist():
    assert set(HANDLERS) == set(get_args(HandlerName))


def test_logs_do_not_contain_request_content(happy, client, caplog):
    caplog.set_level("INFO")
    client.post("/agent", json={"request": "Confidential merger codename ZEBRA-PLAN project plan please."})
    assert "ZEBRA-PLAN" not in caplog.text and "test-key" not in caplog.text


def test_health(client, monkeypatch):
    monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
    assert client.get("/health").json() == {"status": "ok", "llm": "not_configured"}
