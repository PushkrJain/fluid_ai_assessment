import pytest
from docx import Document

from app.document import DocumentSpec, build_docx, heading_matches, missing_sections, verify_docx
from app.schemas import Assumption, DocumentSection, DocumentTable


def spec(**over):
    data = dict(
        run_id="a" * 32, title="Plan", document_type="project plan", audience="Leaders",
        sections=[DocumentSection(heading="Scope", content="One.\n\nTwo **bold**.", bullets=["b1"],
                                  table=DocumentTable(columns=["A", "B"], rows=[["1", "2"]]))],
        assumptions=[Assumption(statement="s", category="illustrative_data", impact="i")],
        open_questions=["q1"], provider="p", model="m",
    )
    data.update(over)
    return DocumentSpec(**data)


def test_build_and_verify_roundtrip(monkeypatch, tmp_path):
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path))
    path = build_docx(spec())
    result = verify_docx(path, "Plan", ["Scope", "Assumptions", "AI-assisted generation disclosure and human-review notice"])
    assert result.ok and result.missing_sections == []
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "**" not in text and "Illustrative data" in " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "Human-Review" in text or "human-review" in text.lower()


def test_missing_heading_is_reported_not_fatal(monkeypatch, tmp_path):
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path))
    result = verify_docx(build_docx(spec()), "Plan", ["Scope", "Budget"])
    assert result.ok and result.missing_sections == ["Budget"]


def test_verification_fails_for_corrupt_missing_or_foreign_files(monkeypatch, tmp_path):
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path))
    corrupt = tmp_path / ("b" * 32 + ".docx")
    corrupt.write_bytes(b"not a docx")
    assert not verify_docx(corrupt, "Plan", []).ok
    assert not verify_docx(tmp_path / ("c" * 32 + ".docx"), "Plan", []).ok
    outside = tmp_path.parent / "x.docx"
    assert not verify_docx(outside, "Plan", []).ok


def test_wrong_title_fails_verification(monkeypatch, tmp_path):
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path))
    assert not verify_docx(build_docx(spec()), "Different Title", []).ok


def test_fallback_notice_is_in_document(monkeypatch, tmp_path):
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path))
    text = " ".join(p.text for p in Document(str(build_docx(spec(fallback_notice="Test notice.")))).paragraphs)
    assert "FALLBACK DOCUMENT" in text


def test_heading_matching_is_tolerant_but_not_loose():
    assert heading_matches("Risks & Mitigations", "Risks and Mitigations")
    assert heading_matches("Scope", "2. Scope")
    assert not heading_matches("Scope", "2. Scope and Exclusions")  # changed on purpose: no substring matching
    assert not heading_matches("Budget", "Timeline")
    assert missing_sections(["Scope", "Budget"], ["Scope and goals"]) == ["Scope", "Budget"]


def build(monkeypatch, tmp_path, sections):
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path))
    return build_docx(spec(sections=sections))


def all_text(path):
    doc = Document(str(path))
    return "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])


def test_markdown_symbols_are_stripped_everywhere(monkeypatch, tmp_path):
    section = DocumentSection(heading="## **Scope**", content="### Heading line\nuse `code` and __bold__",
                              bullets=["**b**"], table=DocumentTable(columns=["**A**"], rows=[["`x`"]]))
    text = all_text(build(monkeypatch, tmp_path, [section]))
    assert not any(token in text for token in ("**", "`", "__", "###"))


def test_numeric_cells_unicode_and_very_long_cells(monkeypatch, tmp_path):
    long_cell = "long " * 2000
    section = DocumentSection(heading="Résumé — 日本語 ✓", content="Ünïcödé ✓ 東京",
                              table=DocumentTable(columns=["N", "Text"], rows=[[1, long_cell], [2.5, "ok"]]))
    text = all_text(build(monkeypatch, tmp_path, [section]))
    assert "日本語" in text and "東京" in text and "2.5" in text and long_cell.strip() in text


def test_empty_table_is_skipped_and_optional_fields_may_be_absent(monkeypatch, tmp_path):
    sections = [DocumentSection(heading="Only a heading"),
                DocumentSection(heading="Empty table", content="c", bullets=[], table=DocumentTable(columns=["A"], rows=[]))]
    path = build(monkeypatch, tmp_path, sections)
    doc = Document(str(path))
    assert len(doc.tables) == 1  # only the assumptions register; the empty table was not rendered
    assert verify_docx(path, "Plan", ["Only a heading", "Empty table"]).missing_sections == []


def test_verify_rejects_an_empty_file(monkeypatch, tmp_path):
    monkeypatch.setenv("OUTPUT_DIR", str(tmp_path))
    empty = tmp_path / ("f" * 32 + ".docx")
    empty.write_bytes(b"")
    assert not verify_docx(empty, "Plan", []).ok


@pytest.mark.parametrize("a, b, expected", [
    ("Stakeholder Analysis", "Stakeholder Analyses", True),   # near-duplicate via the 0.85 ratio, not containment
    ("Project Timeline", "Project Budget", False),            # shared word, different meaning
    ("Scope", "Risks", False), ("Goals", "Costs", False),     # unrelated short headings
    ("3) RISKS & MITIGATIONS", "risks and mitigations", True),  # numbering, case, ampersand
    ("", "Scope", False),
])
def test_heading_matcher_threshold_cases(a, b, expected):
    assert heading_matches(a, b) is expected
